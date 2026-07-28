from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.logger import logger
from home.models import StaffUser
from home.tenancy import get_current_tenant, tenant_qs


def formatINR(number):
    s, *d = str(number).partition(".")
    r = ",".join([s[x - 2:x] for x in range(-3, -len(s), -2)][::-1] + [s[-3:]])
    return "".join([r] + d)


def format_money_compact(value):
    try:
        val = float(value or 0.0)
    except Exception:
        return '0'
    if abs(val - int(val)) < 1e-9:
        return formatINR(format(int(val), '.0f'))
    return formatINR(format(val, '.2f'))


def pdf_base_context(request, applied_filters=None):
    tenant = get_current_tenant(request)
    generated_by = 'System'
    try:
        if request.user.is_authenticated:
            staff = tenant_qs(request, StaffUser).filter(user_ID_id=request.user.pk).first()
            generated_by = staff.name if staff and staff.name else request.user.get_username()
    except Exception:
        generated_by = getattr(request.user, 'username', 'System') or 'System'
    return {
        'tenant': tenant,
        'pdf_brand_name': 'zivabooks.in',
        'pdf_tenant_name': getattr(tenant, 'name', None) or 'zivabooks.in',
        'pdf_generated_by': generated_by,
        'pdf_generated_on': timezone.now(),
        'pdf_applied_filters': applied_filters or [],
    }


def safe_sheet_title(title):
    text = (title or 'Report').strip() or 'Report'
    cleaned = ''.join(ch for ch in text if ch not in '[]:*?/\\')
    return cleaned[:31] or 'Report'


def xlsx_report_response(filename_prefix, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    thin_border = Border(
        left=Side(style='thin', color='D7E0EA'),
        right=Side(style='thin', color='D7E0EA'),
        top=Side(style='thin', color='D7E0EA'),
        bottom=Side(style='thin', color='D7E0EA'),
    )
    wrap_alignment = Alignment(vertical='center', wrap_text=True)
    title_font = Font(name='Calibri', size=15, bold=True)
    subtitle_font = Font(name='Calibri', size=11, color='475569')
    section_font = Font(name='Calibri', size=12, bold=True)
    header_font = Font(name='Calibri', size=11, bold=True)
    meta_font = Font(name='Calibri', size=11, color='334155')
    section_fill = PatternFill('solid', fgColor='E8F1FB')
    header_fill = PatternFill('solid', fgColor='D7E8FB')

    def write_cell(cell, value, font=None, fill=None):
        cell.value = value
        cell.alignment = wrap_alignment
        cell.border = thin_border
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if isinstance(value, (float, Decimal)):
            cell.number_format = '#,##0.00'
        elif isinstance(value, datetime):
            cell.number_format = 'DD-MM-YYYY HH:MM'
        elif hasattr(value, 'strftime') and not isinstance(value, str):
            cell.number_format = 'DD-MM-YYYY'

    for index, spec in enumerate(sheets):
        ws = workbook.create_sheet(title=safe_sheet_title(spec.get('title') or f'Report {index + 1}'))
        row_idx = 1

        def append_row(values, font=None, fill=None):
            nonlocal row_idx
            for col_idx, value in enumerate(values, start=1):
                write_cell(ws.cell(row=row_idx, column=col_idx), value, font=font, fill=fill)
            row_idx += 1

        if spec.get('report_title'):
            append_row([spec['report_title']], font=title_font)
        if spec.get('report_subtitle'):
            append_row([spec['report_subtitle']], font=subtitle_font)
        if spec.get('report_title') or spec.get('report_subtitle'):
            row_idx += 1

        metadata = spec.get('metadata') or []
        if metadata:
            append_row(['Report Metadata'], font=section_font, fill=section_fill)
            for label, value in metadata:
                append_row([label, value], font=meta_font)
            row_idx += 1

        for table in spec.get('tables') or []:
            if table.get('title'):
                append_row([table['title']], font=section_font, fill=section_fill)
            headers = table.get('headers') or []
            if headers:
                append_row(headers, font=header_font, fill=header_fill)
            for row in table.get('rows') or []:
                append_row(row)
            row_idx += 1

        for column_cells in ws.columns:
            length = 0
            letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = '' if cell.value is None else str(cell.value)
                length = max(length, len(value))
            ws.column_dimensions[letter].width = min(max(length + 2, 14), 40)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename={filename_prefix}-{datetime.today().strftime("%Y-%m-%d")}.xlsx'
    response.write(stream.getvalue())
    return response
